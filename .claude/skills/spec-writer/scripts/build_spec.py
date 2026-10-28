#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
시방서 빌더 : content.yaml → 시방서 .docx

사용법
  python build_spec.py <content.yaml> [-o 출력.docx] [--pdf]

- 양식(표지·머리말·안전보건 페이지·Common List·별첨)은 assets/master.docx 를 그대로 사용
- 장 번호는 자동으로 매기며, 본문 속 {sec:키} 를 실제 번호로 바꿈
- [[...]] 는 노란 형광으로 표시되고 마지막에 '확인 필요 목록'으로 출력됨
"""
import argparse
import copy
import re
import shutil
import subprocess
import sys
from pathlib import Path

try:
    import yaml
    from docx import Document
    from docx.shared import Cm
    from lxml import etree
except ImportError as e:  # pragma: no cover
    sys.exit(f"[오류] 필요한 패키지가 없습니다: {e.name}\n  pip install python-docx pyyaml lxml")

SKILL_DIR = Path(__file__).resolve().parent.parent
MASTER = SKILL_DIR / "assets" / "master.docx"
DATA = SKILL_DIR / "data"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

SECTION_KEYS = ["name", "due", "location", "spec", "env", "scope", "detail",
                "install", "misc", "docs", "warranty", "common"]

warnings: list[str] = []
checks: list[str] = []


# ───────────────────────── 기본 XML 도우미 ─────────────────────────
def el(tag, attrs=None, children=()):
    e = etree.Element(W + tag)
    for k, v in (attrs or {}).items():
        e.set(W + k, str(v))
    for c in children:
        e.append(c)
    return e


def base_rpr(bold=False, size=None):
    rpr = el("rPr")
    if bold:
        rpr.append(el("b"))
    rpr.append(el("color", {"val": "000000", "themeColor": "text1"}))
    if size:
        rpr.append(el("sz", {"val": size}))
    rpr.append(el("szCs", {"val": size or 18}))
    return rpr


INLINE_RE = re.compile(r"(\*\*.+?\*\*|\[\[.+?\]\]|\{\{red\}\}.+?\{\{/red\}\})")


def add_runs(p, text, rpr_proto):
    """text 를 run 으로 추가. **굵게**, [[확인 필요]], {{red}}빨강{{/red}}, 줄바꿈(\\n) 지원"""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        rpr = copy.deepcopy(rpr_proto)
        if part.startswith("**") and part.endswith("**"):
            part = part[2:-2]
            if rpr.find(W + "b") is None:
                rpr.insert(0, el("b"))
        elif part.startswith("[[") and part.endswith("]]"):
            part = part[2:-2]
            checks.append(part)
            rpr.append(el("highlight", {"val": "yellow"}))
        elif part.startswith("{{red}}"):
            part = part[len("{{red}}"):-len("{{/red}}")]
            for c in rpr.findall(W + "color"):
                rpr.remove(c)
            rpr.append(el("color", {"val": "FF0000"}))
        lines = part.split("\n")
        for i, line in enumerate(lines):
            r = el("r")
            r.append(copy.deepcopy(rpr))
            if i > 0:
                r.append(el("br"))
            t = el("t")
            t.text = line
            t.set(XML_SPACE, "preserve")
            r.append(t)
            p.append(r)


def para(text="", left=None, hanging=None, first_line=None, bold=False,
         keep_next=False, before=None, jc=None, size=None):
    p = el("p")
    ppr = el("pPr")
    if keep_next:
        ppr.append(el("keepNext"))
    sp = {"after": 0}
    if before:
        sp["before"] = before
    ppr.append(el("spacing", sp))
    ind = {}
    if left is not None:
        ind["left"] = left
    if hanging is not None:
        ind["hanging"] = hanging
    if first_line is not None:
        ind["firstLine"] = first_line
    if ind:
        ppr.append(el("ind", ind))
    if jc:
        ppr.append(el("jc", {"val": jc}))
    ppr.append(base_rpr(bold, size))
    p.append(ppr)
    if text:
        add_runs(p, text, base_rpr(bold, size))
    return p


def spacer(size=8):
    return para("", size=size)


# ───────────────────────── 블록 렌더링 ─────────────────────────
NUM_ITEM_RE = re.compile(r"^\d+\)\s")
CIRCLE_RE = re.compile(r"^[①-⑳]")
H3_NUM_RE = re.compile(r"^\d+\)\s*")
H4_NUM_RE = re.compile(r"^\d+-\d+\s")


class Ctx:
    """세부 사양 블록 렌더링 상태 (들여쓰기 기준, 자동 번호)"""

    def __init__(self, base=800, h3_start=1):
        self.base = base
        self.h3 = h3_start - 1
        self.h4 = 0
        self.table_ind = 411


def render_blocks(blocks, ctx, src_dir):
    out = []
    for b in blocks or []:
        if isinstance(b, dict):
            if "table" in b:
                out += render_table(b["table"], ctx)
            elif "image" in b:
                out.append(("IMAGE", b, src_dir))
            elif "blocks" in b:  # 묶음
                out += render_blocks(b["blocks"], ctx, src_dir)
            else:
                warnings.append(f"알 수 없는 블록: {b}")
            continue
        s = str(b).rstrip()
        if s == "":
            out.append(spacer())
        elif s.startswith("## "):  # 소소제목 3-1
            t = s[3:].strip()
            ctx.h4 += 1
            if not H4_NUM_RE.match(t):
                t = f"{ctx.h3}-{ctx.h4} {t}"
            out.append(para(t, left=960, bold=True, keep_next=True, before=40))
            ctx.base = 1160
        elif s.startswith("# "):  # 소제목 1)
            t = s[2:].strip()
            ctx.h3 += 1
            ctx.h4 = 0
            if not H3_NUM_RE.match(t):
                t = f"{ctx.h3}) {t}"
            out.append(para(t, left=760, bold=True, keep_next=True, before=80))
            ctx.base = 1160
        elif s.startswith("- ") or s == "-":
            out.append(para(s, left=ctx.base, hanging=200))
        elif s.startswith("-. "):
            out.append(para(s, left=ctx.base, hanging=280))
        elif s.startswith(": "):
            out.append(para(s, left=ctx.base + 240, hanging=200))
        elif s.startswith("→"):
            out.append(para(s, left=ctx.base + 240, hanging=260))
        elif s.startswith("※") or s.startswith("* "):
            out.append(para(s, left=ctx.base, hanging=200))
        elif CIRCLE_RE.match(s):
            out.append(para(s, left=ctx.base + 240, hanging=300))
        elif NUM_ITEM_RE.match(s):
            out.append(para(s, left=ctx.base + 300, hanging=300))
        else:
            out.append(para(s, left=ctx.base))
    return out


def render_table(spec, ctx, total=9639):
    header = spec.get("header") or []
    rows = spec.get("rows") or []
    ncol = len(header) if header else max(len(r) for r in rows)
    for r in rows:
        if len(r) > ncol:
            warnings.append(f"표 행의 칸 수({len(r)})가 열 수({ncol})보다 많음 — 셀 안 쉼표 때문일 수 있으니 따옴표로 감싸세요: {r}")
    widths = spec.get("widths") or [1] * ncol
    if len(widths) != ncol:
        warnings.append(f"표 widths 개수({len(widths)})가 열 수({ncol})와 다름: {header}")
        widths = [1] * ncol
    s = float(sum(widths))
    tw = [int(total * w / s) for w in widths]
    tw[-1] += total - sum(tw)
    align = spec.get("align") or (["c"] + ["l"] * (ncol - 1))
    jcmap = {"c": "center", "l": "left", "r": "right"}

    tbl = el("tbl")
    tblpr = el("tblPr")
    tblpr.append(el("tblStyle", {"val": "a6"}))
    tblpr.append(el("tblW", {"w": 0, "type": "auto"}))
    tblpr.append(el("tblInd", {"w": ctx.table_ind, "type": "dxa"}))
    borders = el("tblBorders")
    for side in ("top", "left", "bottom", "right"):
        borders.append(el(side, {"val": "single", "sz": 12, "space": 0, "color": "auto"}))
    tblpr.append(borders)
    tblpr.append(el("tblLook", {"val": "04A0", "firstRow": 1, "lastRow": 0, "firstColumn": 1,
                                "lastColumn": 0, "noHBand": 0, "noVBand": 1}))
    tbl.append(tblpr)
    grid = el("tblGrid")
    for w in tw:
        grid.append(el("gridCol", {"w": w}))
    tbl.append(grid)

    def row(cells, is_head):
        tr = el("tr")
        trpr = el("trPr")
        trpr.append(el("cantSplit"))
        if is_head:
            trpr.append(el("tblHeader"))
        tr.append(trpr)
        cells = list(cells) + [""] * (ncol - len(cells))
        for i, c in enumerate(cells[:ncol]):
            tc = el("tc")
            tcpr = el("tcPr")
            tcpr.append(el("tcW", {"w": tw[i], "type": "dxa"}))
            text = "" if c is None else str(c)
            if text == "^":
                tcpr.append(el("vMerge"))
                text = ""
            elif not is_head and i < ncol and _next_is_merge(rows, cells, i):
                tcpr.append(el("vMerge", {"val": "restart"}))
            if is_head:
                tcpr.append(el("shd", {"val": "clear", "color": "auto", "fill": "F2F2F2"}))
            tcpr.append(el("vAlign", {"val": "center"}))
            tc.append(tcpr)
            jc = "center" if is_head else jcmap.get(align[i] if i < len(align) else "l", "left")
            p = para("", jc=jc, bold=is_head)
            p.find(W + "pPr").find(W + "spacing").set(W + "line", "240")
            p.find(W + "pPr").find(W + "spacing").set(W + "lineRule", "auto")
            if text:
                add_runs(p, text, base_rpr(is_head))
            tc.append(p)
            tr.append(tc)
        return tr

    _MERGE_STATE["rows"] = rows
    if header:
        tbl.append(row(header, True))
    for ri, r in enumerate(rows):
        _MERGE_STATE["ri"] = ri
        tbl.append(row(r, False))
    return [keep_table_together(tbl), spacer()]


_MERGE_STATE = {"rows": [], "ri": 0}


def _next_is_merge(rows, cells, col):
    ri = _MERGE_STATE["ri"]
    if ri + 1 < len(rows):
        nxt = rows[ri + 1]
        return col < len(nxt) and str(nxt[col]) == "^"
    return False


def keep_table_together(tbl, max_rows=15):
    """행이 적은 표는 페이지에서 쪼개지지 않도록 마지막 행 전까지 '다음 문단과 함께'"""
    rows = tbl.findall(W + "tr")
    if len(rows) > max_rows:
        return tbl
    for tr in rows[:-1]:
        for p in tr.iter(W + "p"):
            ppr = p.find(W + "pPr")
            if ppr is None:
                ppr = el("pPr")
                p.insert(0, ppr)
            if ppr.find(W + "keepNext") is None:
                ppr.insert(0 if ppr.find(W + "pStyle") is None else 1, el("keepNext"))
    return tbl


# ───────────────────────── 원형 표 채우기 (Pipe Size / 공급 범위) ─────────────────────────
def set_cell(tc, text):
    ps = tc.findall(W + "p")
    p = ps[0]
    for extra in ps[1:]:
        tc.remove(extra)
    proto = None
    r0 = p.find(W + "r")
    if r0 is not None and r0.find(W + "rPr") is not None:
        proto = copy.deepcopy(r0.find(W + "rPr"))
    for ch in list(p):
        if ch.tag != W + "pPr":
            p.remove(ch)
    if proto is None:
        ppr = p.find(W + "pPr")
        mark = ppr.find(W + "rPr") if ppr is not None else None
        proto = copy.deepcopy(mark) if mark is not None else base_rpr()
    for b in proto.findall(W + "b"):
        proto.remove(b)
    jc = p.find(f"{W}pPr/{W}jc")
    if jc is not None and jc.get(W + "val") == "both":
        jc.set(W + "val", "left")
    if text:
        add_runs(p, str(text), proto)


def fill_pipe_table(tbl, pipe):
    rows = tbl.findall(W + "tr")
    # 행 순서 : 외 경 / 두 께 / 길 이
    labels = ["외 경", "두 께", "길 이"]
    for r, label in zip(rows[1:4], labels):
        vals = pipe.get(label) or pipe.get(label.replace(" ", "")) or ["", ""]
        vals = list(vals) + ["", "", ""]
        tcs = r.findall(W + "tc")
        for i in range(3):
            set_cell(tcs[i + 1], vals[i])
        if not any(vals[:2]):
            checks.append(f"제품사양 Pipe Size '{label}' 값 없음")
    return keep_table_together(tbl)


def fill_scope_table(tbl, items):
    rows = tbl.findall(W + "tr")
    proto = rows[2]
    for r in rows[2:]:
        tbl.remove(r)
    for i, it in enumerate(items, 1):
        it = list(it) + ["", "", "", ""]
        tr = copy.deepcopy(proto)
        tcs = tr.findall(W + "tc")
        vals = [str(i), it[0], it[1], it[2], it[3]]
        for tc, v in zip(tcs, vals):
            set_cell(tc, v)
        tbl.append(tr)
    return keep_table_together(tbl)


# ───────────────────────── 메인 조립 ─────────────────────────
def load_yaml(p):
    with open(p, encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def leaf_paras(root):
    for p in root.iter(W + "p"):
        if p.find(".//" + W + "p") is None:
            yield p


def ptext(p):
    return "".join(t.text or "" for t in p.iter(W + "t"))


def replace_placeholders(root, mapping):
    for t in root.iter(W + "t"):
        if t.text and "{{" in t.text:
            for k, v in mapping.items():
                t.text = t.text.replace("{{%s}}" % k, v)


def resolve_tokens(obj, secmap, extra):
    """중첩 구조 안의 모든 문자열에서 {sec:키}, {warranty_site} 치환"""
    if isinstance(obj, str):
        def rep(m):
            key = m.group(1)
            if key not in secmap:
                warnings.append(f"알 수 없는 장 참조 {{sec:{key}}} (사용 가능: {', '.join(secmap)})")
                return m.group(0)
            return str(secmap[key])
        s = re.sub(r"\{sec:(\w+)\}", rep, obj)
        for k, v in extra.items():
            s = s.replace("{%s}" % k, v)
        return s
    if isinstance(obj, list):
        return [resolve_tokens(x, secmap, extra) for x in obj]
    if isinstance(obj, dict):
        return {k: resolve_tokens(v, secmap, extra) for k, v in obj.items()}
    return obj


def h1(num, title):
    return para(f"{num}. {title}", first_line=400, bold=True, keep_next=True, before=120)


def h1_label(num, label, value):
    p = para("", first_line=400, before=60)
    add_runs(p, f"{num}. **{label}** : {value}", base_rpr())
    return p


def h2(num, title):
    return para(f"{num} {title}", first_line=600, bold=True, keep_next=True, before=120)


def build(content_path: Path, out_path: Path):
    c = load_yaml(content_path)
    src_dir = content_path.parent
    factories = load_yaml(DATA / "factories.yaml")
    bp = load_yaml(DATA / "boilerplate.yaml")
    cfg = load_yaml(DATA / "config.yaml")

    for req in ("title", "factory"):
        if not c.get(req):
            sys.exit(f"[오류] content.yaml 에 '{req}' 가 없습니다.")
    fac = factories.get(c["factory"])
    if fac is None:
        sys.exit(f"[오류] factories.yaml 에 '{c['factory']}' 가 없습니다. (있는 공장: {', '.join(factories)})")

    # 제품사양 Pipe Size
    pipe = c.get("pipe_size")
    if pipe is None:
        line = c.get("line")
        lines = fac.get("lines") or {}
        if line not in lines:
            sys.exit(f"[오류] {c['factory']} 에 라인 '{line}' 이 없습니다. (있는 라인: {', '.join(lines)})"
                     "\n  factories.yaml 에 추가하거나 content.yaml 에 pipe_size 를 직접 쓰세요.")
        pipe = lines[line]["pipe_size"]

    if len(c.get("cover_title", c["title"])) > 40:
        warnings.append("제목이 40자를 넘어 머리말에서 두 줄이 될 수 있습니다 (cover_title 로 줄여 쓸 수 있음)")

    # 장 번호 매기기
    order = ["name", "due"]
    if c.get("location"):
        order.append("location")
    order += ["spec", "env", "scope", "detail", "install", "misc", "docs", "warranty", "common"]
    secmap = {k: i for i, k in enumerate(order, 1)}
    extra = {"warranty_site": fac.get("warranty_site", "")}
    c = resolve_tokens(c, secmap, extra)
    bp = resolve_tokens(bp, secmap, extra)

    doc = Document(str(MASTER))
    body = doc.element.body

    # 원형 표 꺼내기
    marker = holder_scope = holder_pipe = None
    for p in body.findall(W + "p"):
        tx = ptext(p)
        if tx == "{{BODY}}":
            marker = p
        elif tx == "{{PROTO_SCOPE}}":
            holder_scope = p
        elif tx == "{{PROTO_PIPE}}":
            holder_pipe = p
    proto_scope = holder_scope.getnext()
    proto_pipe = holder_pipe.getnext()

    out = []  # (lxml element) 또는 ("IMAGE", spec, dir)

    # 1~3. 품명 / 납기 / 설치 장소
    out.append(h1_label(secmap["name"], "품  명", c["title"]))
    out.append(h1_label(secmap["due"], c.get("due_label", "제작 설치 및 시운전 납기"),
                        f"{c.get('due', '[[납기]]')} (납기는 추후 변경 가능)"))
    if c.get("location"):
        out.append(h1_label(secmap["location"], "설치 장소", c["location"]))

    # 제품사양
    out.append(h1(secmap["spec"], "제품사양"))
    out.append(para("1) Pipe Size", left=760, bold=True, keep_next=True, before=40))
    out.append(fill_pipe_table(copy.deepcopy(proto_pipe), pipe))
    out.append(spacer())
    ctx = Ctx(base=1160)
    for i, sub in enumerate(c.get("spec_extra") or [], 2):
        out.append(para(f"{i}) {sub['title']}", left=760, bold=True, keep_next=True, before=80))
        out += render_blocks(sub.get("blocks"), ctx, src_dir)

    # 설치환경 (고정)
    out.append(h1(secmap["env"], "설치환경"))
    for line in bp["env"]:
        out.append(para(line, left=760))

    # 공급 범위
    out.append(h1(secmap["scope"], "공급 범위"))
    out.append(para("1) 공급 범위 구분", left=760, bold=True, keep_next=True, before=40))
    out.append(fill_scope_table(copy.deepcopy(proto_scope), c.get("scope") or bp["scope"]))
    out.append(spacer())
    out.append(para("2) 공급 List", left=760, bold=True, keep_next=True, before=80))
    supply = c.get("supply_list") or []
    for r in supply:
        if len(r) > 3:
            warnings.append(f"supply_list 행이 3칸(세부 내용, 수량, 비고)보다 많음 — 따옴표로 감싸세요: {r}")
    for r in c.get("scope") or []:
        if len(r) > 4:
            warnings.append(f"scope 행이 4칸(항목, 세아제강, 수주처, 비고)보다 많음 — 따옴표로 감싸세요: {r}")
    if not supply:
        warnings.append("supply_list(공급 List)가 비어 있습니다.")
    out += render_table({
        "header": ["No.", "세부 내용", "수 량", "비 고"],
        "widths": [8, 52, 12, 28],
        "align": ["c", "l", "c", "l"],
        "rows": [[i] + list(r) for i, r in enumerate(supply, 1)],
    }, Ctx())

    # 공급 설비별 세부 사양
    n = secmap["detail"]
    out.append(h1(n, f"공급 설비별 세부 사양 {c.get('detail_note', bp['detail_note'])}".strip()))
    details = list(c.get("details") or [])
    if not details:
        warnings.append("details(공급 설비별 세부 사양)가 비어 있습니다.")
    if c.get("safety_permits", True):
        details.append(bp["safety_permits"])
    for i, sec in enumerate(details, 1):
        out.append(h2(f"{n}-{i}", sec["title"]))
        out += render_blocks(sec.get("blocks"), Ctx(base=800), src_dir)
        out.append(spacer())

    # 설치 공사 및 시운전
    n = secmap["install"]
    out.append(h1(n, c.get("install_title", "설치 공사 및 시운전")))
    for i, sec in enumerate(c.get("install") or bp["install"], 1):
        out.append(h2(f"{n}-{i}", sec["title"]))
        out += render_blocks(sec.get("blocks"), Ctx(base=800), src_dir)
    out.append(spacer())

    # 기타사항 / 문서
    for key, title in (("misc", "기타사항"), ("docs", "문 서")):
        n = secmap[key]
        out.append(h1(n, title))
        items = list(c.get(key) or bp[key]) + list(c.get(f"{key}_extra") or [])
        for i, it in enumerate(items, 1):
            out.append(para(f"{n}-{i}. {it}", left=1000, hanging=400))
        out.append(spacer())

    # Warranty (고정)
    n = secmap["warranty"]
    out.append(h1(n, "Warranty"))
    for i, sec in enumerate(bp["warranty"], 1):
        out.append(para(f"{n}-{i}. {sec['title']}", left=600, before=60))
        out += render_blocks(sec["blocks"], Ctx(base=600), src_dir)
    out.append(spacer())

    # Common List 제목 (표는 양식에 있음)
    out.append(para(f"{secmap['common']}. Common List", first_line=400, bold=True, keep_next=True))

    # 삽입
    for item in out:
        if isinstance(item, tuple) and item[0] == "IMAGE":
            _, spec, d = item
            path = (d / spec["image"]).resolve()
            if not path.exists():
                warnings.append(f"그림 파일 없음: {path}")
                p = para("", jc="center", before=60)
                add_runs(p, f"[[그림 없음: {spec['image']}]]", base_rpr())
                marker.addprevious(p)
            else:
                dp = doc.add_paragraph()
                dp.alignment = 1  # center
                dp.paragraph_format.space_after = 0
                dp.add_run().add_picture(str(path), width=Cm(float(spec.get("width_cm", 12))))
                marker.addprevious(dp._p)
            if spec.get("caption"):
                marker.addprevious(para(spec["caption"], jc="center"))
        else:
            marker.addprevious(item)
    for e in (marker, holder_scope, proto_scope, holder_pipe, proto_pipe):
        body.remove(e)

    # 표지 / 머리말 / 담당자
    a = dict(cfg.get("author") or {})
    a.update(c.get("author") or {})
    mapping = {
        "COMPANY": fac.get("company", "㈜세아제강"),
        "FACTORY": c["factory"],
        "TITLE": c.get("cover_title", c["title"]),
        "AUTHOR_LINE": f"{c['factory']} {a.get('team', '')} {a.get('name', '')} {a.get('rank', '')}, Tel : {a.get('tel', '')}",
    }
    replace_placeholders(body, mapping)
    for s in doc.sections:
        replace_placeholders(s.header._element, mapping)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))

    # 남은 자리표시자 검사
    left = [ptext(p) for p in leaf_paras(body) if "{{" in ptext(p) or "{sec:" in ptext(p)]
    for s in doc.sections:
        left += [ptext(p) for p in leaf_paras(s.header._element) if "{{" in ptext(p)]
    for t in left:
        warnings.append(f"치환 안 된 자리표시자: {t[:60]}")
    return secmap


def to_pdf(docx: Path):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        for cand in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                     r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
            if Path(cand).exists():
                exe = cand
                break
    if not exe:
        print("  (PDF 변환 생략: LibreOffice 없음)")
        return None
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    pdf = docx.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="content.yaml → 시방서 docx")
    ap.add_argument("content")
    ap.add_argument("-o", "--out")
    ap.add_argument("--pdf", action="store_true", help="LibreOffice 로 PDF 도 생성")
    args = ap.parse_args()
    cp = Path(args.content).resolve()
    c = load_yaml(cp)
    out = Path(args.out) if args.out else cp.parent / f"{c.get('file_prefix', '')}{c.get('title', 'spec')}_시방서.docx"
    secmap = build(cp, out.resolve())
    print(f"[완료] {out}")
    print("  장 구성 : " + " / ".join(f"{v}.{k}" for k, v in secmap.items()))
    if args.pdf:
        pdf = to_pdf(out.resolve())
        if pdf:
            print(f"  PDF : {pdf}")
    if checks:
        print(f"\n[확인 필요 {len(checks)}건] (문서에 노란 형광 표시)")
        for x in dict.fromkeys(checks):
            print("  - " + x)
    if warnings:
        print(f"\n[경고 {len(warnings)}건]")
        for w in dict.fromkeys(warnings):
            print("  - " + w)
        sys.exit(2 if any(k in w for w in warnings for k in ("치환 안 된", "알 수 없는 장 참조", "따옴표")) else 0)


if __name__ == "__main__":
    main()
